from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('/Users/rohan/Documents/ChatGPT/Linear-clone/Scoutline_Product_Requirements.docx')

# compact_reference_guide preset
BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
INK = '17202A'
MUTED = '5B6570'
LIGHT_FILL = 'E8EEF5'
PALE_FILL = 'F4F6F9'
BORDER = 'B8C3CF'
WHITE = 'FFFFFF'


def set_run_font(run, name='Calibri', size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)
    shd.set(qn('w:val'), 'clear')


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_border(cell, color=BORDER, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = borders.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_table_geometry(table, widths, indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent))
    tbl_ind.set(qn('w:type'), 'dxa')
    layout = tbl_pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tbl_pr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(widths[index]))
            tc_w.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    separate.append(text)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)
    set_run_font(run, size=9, color=MUTED)


def add_numbering(doc, kind='bullet'):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn('w:abstractNumId'))) for node in numbering.findall(qn('w:abstractNum'))]
    num_ids = [int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    multi = OxmlElement('w:multiLevelType')
    multi.set(qn('w:val'), 'singleLevel')
    abstract.append(multi)
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:start')
    start.set(qn('w:val'), '1')
    level.append(start)
    fmt = OxmlElement('w:numFmt')
    fmt.set(qn('w:val'), 'bullet' if kind == 'bullet' else 'decimal')
    level.append(fmt)
    text = OxmlElement('w:lvlText')
    text.set(qn('w:val'), '●' if kind == 'bullet' else '%1.')
    level.append(text)
    jc = OxmlElement('w:lvlJc')
    jc.set(qn('w:val'), 'left')
    level.append(jc)
    ppr = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'num')
    tab.set(qn('w:pos'), '540')
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '540')
    ind.set(qn('w:hanging'), '270')
    ppr.append(ind)
    level.append(ppr)
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Calibri')
    fonts.set(qn('w:hAnsi'), 'Calibri')
    rpr.append(fonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    abstract_ref = OxmlElement('w:abstractNumId')
    abstract_ref.set(qn('w:val'), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn('w:numPr'))
    if num_pr is None:
        num_pr = OxmlElement('w:numPr')
        ppr.append(num_pr)
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    num = OxmlElement('w:numId')
    num.set(qn('w:val'), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def paragraph_format(paragraph, before=0, after=6, line=1.25):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_para(doc, text='', bold_prefix=None, italic=False, after=6, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph_format(paragraph, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(paragraph.add_run(bold_prefix), bold=True)
        set_run_font(paragraph.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_run_font(paragraph.add_run(text), italic=italic)
    return paragraph


def add_bullet(doc, text, num_id):
    paragraph = doc.add_paragraph()
    paragraph_format(paragraph, after=4)
    apply_num(paragraph, num_id)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_numbered(doc, text, num_id):
    paragraph = doc.add_paragraph()
    paragraph_format(paragraph, after=4)
    apply_num(paragraph, num_id)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9120], indent=120)
    # Keep a callout intact. Without this, Word may split the final line of a
    # short note onto the next page, which makes the requirement harder to read.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_FILL)
    set_cell_border(cell, color='C8D2DD', size='6')
    p = cell.paragraphs[0]
    paragraph_format(p, after=0, line=1.2)
    set_run_font(p.add_run(f'{label}: '), size=10, color=DARK_BLUE, bold=True)
    set_run_font(p.add_run(text), size=10, color=INK)
    spacer = doc.add_paragraph()
    paragraph_format(spacer, after=2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths, indent=120)
    table.style = 'Table Grid'
    # Tell assistive technology and repeated-page rendering that this is the
    # table's header row. (Single-cell callouts use add_note and are not data
    # tables, so they intentionally do not receive this marker.)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:tblHeader'))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        paragraph_format(p, after=0, line=1.15)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(header), size=9.5, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            set_cell_border(cell, color='D5DDE5', size='4')
            p = cell.paragraphs[0]
            paragraph_format(p, after=0, line=1.15)
            set_run_font(p.add_run(str(value)), size=9.5, color=INK)
    # Re-apply geometry after adding data rows so every row has explicit widths,
    # cell margins, and fixed layout (not only the initial header row).
    set_table_geometry(table, widths, indent=120)
    spacer = doc.add_paragraph()
    paragraph_format(spacer, after=2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f'Heading {level}')
    paragraph_format(paragraph, after={1: 10, 2: 7, 3: 5}[level], line=1.15)
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return paragraph


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def set_document_defaults(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ('Heading 1', 16, BLUE, 18, 10),
        ('Heading 2', 13, BLUE, 14, 7),
        ('Heading 3', 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    # Quiet running header and page-number footer.
    header = section.header
    hp = header.paragraphs[0]
    paragraph_format(hp, after=0, line=1.0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(hp.add_run('Scoutline · Product requirements and workflow'), size=8.5, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    paragraph_format(fp, after=0, line=1.0)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(fp.add_run('Scoutline · '), size=8.5, color=MUTED)
    add_field(fp, 'PAGE')


def build_document():
    doc = Document()
    set_document_defaults(doc)
    bullet_id = add_numbering(doc, 'bullet')
    decimal_id = add_numbering(doc, 'decimal')

    # First-page masthead.
    spacer = doc.add_paragraph()
    paragraph_format(spacer, after=4)
    title = doc.add_paragraph()
    paragraph_format(title, after=4, line=1.0)
    set_run_font(title.add_run('SCOUTLINE'), size=23, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    paragraph_format(subtitle, after=16, line=1.1)
    set_run_font(subtitle.add_run('Product requirements and workflow description'), size=15, color=MUTED)
    meta_rows = [
        ('Purpose', 'A design-handoff description of the product’s current behavior and requirements.'),
        ('Audience', 'Product, UX, UI, and workflow collaborators.'),
        ('Scope', 'Current product capabilities; excludes visual theme, color, layout, and implementation details.'),
        ('Product status', 'Local-first MVP in active development.'),
        ('Last updated', 'August 21, 2026'),
    ]
    add_table(doc, ['Document field', 'Definition'], meta_rows, [1900, 7220])
    add_note(doc, 'Reading guide', 'Treat this document as the functional source of truth for a design exploration. It describes what the product must help a person understand and do, not how the interface should look or how the services are implemented.')

    add_heading(doc, '1. Product at a glance', 1)
    add_para(doc, 'Scoutline is a local-first lead research and sales workspace for finding service businesses that may benefit from a better website, understanding the opportunity with evidence, preparing a relevant first message, and manually moving the relationship through sales to a contract.')
    add_para(doc, 'The product is an AI-assisted sales researcher, not an autonomous sales agent. Its job is to reduce the time between “this may be a good prospect” and “I understand why this prospect is worth contacting.” The user remains responsible for deciding who is qualified, reviewing outreach, sending it, and managing the relationship.')
    add_heading(doc, 'Primary job to be done', 2)
    add_para(doc, 'Given a geography and a set of business categories, help the user discover businesses, collect a usable profile, inspect the public website, identify meaningful opportunities, and turn the strongest findings into a concise, evidence-backed outreach conversation.')
    add_heading(doc, 'The product promise', 2)
    add_bullet(doc, 'Every important conclusion should be traceable to a source, observation, or saved piece of evidence.', bullet_id)
    add_bullet(doc, 'Research and AI reasoning should clarify a human decision; they should not silently make the sales decision.', bullet_id)
    add_bullet(doc, 'Unknown information remains unknown instead of being guessed.', bullet_id)
    add_bullet(doc, 'Outreach is prepared for review, never sent automatically.', bullet_id)
    add_bullet(doc, 'Repeated collection and analysis should remain understandable: the user can tell what was discovered in which run and what changed over time.', bullet_id)
    add_heading(doc, 'Primary user', 2)
    add_para(doc, 'A founder, independent studio, or small sales/research team looking for local service-business clients—especially businesses where a website improvement could plausibly improve calls, quote requests, bookings, or trust.')

    add_heading(doc, '2. Product principles and guardrails', 1)
    principles = [
        ('Evidence before opinion', 'Scores, findings, and outreach angles should be grounded in observable business or website evidence.'),
        ('Deterministic score, advisory AI', 'The opportunity score is calculated from measured signals. Muse explains and synthesizes; it does not replace or silently alter the score.'),
        ('Human approval at the consequential steps', 'The user decides which clients become qualified, which email version to use, when to mark a relationship stage, and when a contract is recorded.'),
        ('Truthful uncertainty', 'Revenue, profit, employee count, branches, decision makers, and similar metrics are shown as unknown unless a named source supplies them.'),
        ('Traceability over novelty', 'A user should be able to answer: where did this client come from, when was it collected, what was observed, and which run produced the result?'),
        ('Operational failure is not rejection', 'A timeout, blocked crawl, provider failure, or Muse failure should be represented as a recoverable operational state, not as a negative judgment about the business.'),
        ('No spam machine', 'The system supports draft creation and manual review. It does not automatically send cold emails or continue after an opt-out.'),
    ]
    add_table(doc, ['Principle', 'Requirement'], principles, [2300, 6820])

    add_heading(doc, '3. End-to-end product flow', 1)
    add_para(doc, 'The core workflow moves a business from a raw discovery record to a human-reviewed sales opportunity. The user can enter through live discovery or a CSV import, and can also run individual actions for an existing client.')
    flow_steps = [
        ('Start or import data', 'Choose a location, business categories, and a maximum client count for a collection run, or import a prepared list of businesses.'),
        ('Collect business records', 'Collect public business information such as name, category, address, website, contact paths, ratings, reviews, and provider-supplied business signals.'),
        ('Normalize and identify provenance', 'Deduplicate businesses primarily by website/domain and record whether each client was new to the workspace, matched from a previous run, or audited again.'),
        ('Analyze websites', 'Run a bounded website audit that combines HTML observations, limited crawling, internal-link checks, and Lighthouse measurements where available.'),
        ('Prepare a recommendation', 'Save the measured opportunity score and evidence. Optionally ask Muse to synthesize the evidence into priorities, recommendations, strengths, unknowns, and qualification guidance.'),
        ('Qualify manually', 'The user decides whether the client is worth pursuing. Qualifying a client adds it to the sales pipeline.'),
        ('Prepare outreach', 'Generate three materially different email lengths from the saved findings, then edit, copy, and review the selected version before sending it manually.'),
        ('Manage the relationship', 'Move the client through the sales stages. If the client is won, record the agreed amount, scope, terms, dates, conditions, and next step.'),
    ]
    for label, detail in flow_steps:
        p = add_numbered(doc, '', decimal_id)
        set_run_font(p.add_run(label + ' — '), bold=True)
        set_run_font(p.add_run(detail))
    add_note(doc, 'Important vocabulary decision', '“Approved” is not a research checkpoint. “Qualified” is the user’s sales decision and should place the client in the pipeline. “Ranked” is not a separate stage; the opportunity score is a prioritization signal. “LLM insights” is not a separate checkpoint; Muse is an optional advisory layer inside the analyzed/recommendation experience.')

    add_heading(doc, '4. Core domain concepts', 1)
    domain_rows = [
        ('Client', 'A business record in the workspace. It contains identity, contact paths, source information, research state, findings, sales state, notes, and (when won) contract details.'),
        ('Research run', 'A batch-oriented collection or audit session defined by location, categories, source, time, and limits. A run has its own status, progress, counts, and linked clients.'),
        ('Finding / evidence', 'An observed website or business fact, such as a Lighthouse metric, page structure observation, broken-link result, form friction, CTA issue, technology signal, review signal, or source-backed profile field.'),
        ('Opportunity score', 'A deterministic 0–100 signal used to prioritize review. It is separate from confidence and is not a guarantee of commercial value.'),
        ('Confidence', 'A measure of how complete and reliable the evidence is. A high opportunity score with low confidence should invite investigation rather than certainty.'),
        ('Muse synthesis', 'An optional evidence-grounded explanation that summarizes the most important opportunities, recommendations, strengths, unknowns, and qualification guidance. It does not change the deterministic score.'),
        ('Outreach draft', 'A generated or fallback email subject/body tied to the client’s findings. Three length levels are available and each can be edited before copying.'),
        ('Sales stage', 'The manual relationship state: Qualified, Contacted, Replied, Meeting, Proposal, Won, or Lost.'),
        ('Contract record', 'A manual post-win record containing amount, currency, scope, payment terms, signed date, start date, conditions, and next step.'),
        ('Run provenance', 'The relationship between a client and one or more runs: added as new, matched in a later discovery, or audited by that run.'),
    ]
    add_table(doc, ['Concept', 'Meaning in the product'], domain_rows, [2100, 7020])

    add_heading(doc, '5. Research state model', 1)
    add_para(doc, 'Research progress answers “how much do we know about this client?” It is deliberately separate from the sales pipeline, which answers “where are we in the relationship?”')
    state_rows = [
        ('Scraped', 'A client record and public business data exist in the workspace.', 'The business has been collected or imported. Website analysis may still be pending.'),
        ('Analyzed', 'Website evidence and the opportunity score are saved.', 'The deterministic audit has completed, even if one component such as Lighthouse is partial or failed.'),
        ('Recommended', 'A recommendation is ready for the user’s decision.', 'This can use the completed deterministic analysis and may include Muse’s optional advisory synthesis. It does not mean the user has qualified the client.'),
    ]
    add_table(doc, ['Research phase', 'Definition of done', 'What it does not mean'], state_rows, [1700, 3200, 4460])
    add_heading(doc, 'Operational statuses', 2)
    add_para(doc, 'Any research action can also carry an operational status independent of the research phase: Queued, Running, Waiting, Completed, Failed, Paused, or Cancelled. A failure is attached to a stage and includes a reason and retry path.')
    add_bullet(doc, 'Queued: work has been accepted but has not started.', bullet_id)
    add_bullet(doc, 'Running: work is actively being performed.', bullet_id)
    add_bullet(doc, 'Waiting: the system is waiting on an external provider or a retry window.', bullet_id)
    add_bullet(doc, 'Completed: the requested operation finished and its result is saved.', bullet_id)
    add_bullet(doc, 'Failed: the operation could not finish; the reason is visible and retryable where appropriate.', bullet_id)
    add_bullet(doc, 'Paused or Cancelled: a run was deliberately stopped by the user.', bullet_id)

    add_heading(doc, '6. Business discovery and client data', 1)
    add_heading(doc, 'Ways a client enters the workspace', 2)
    add_bullet(doc, 'Live discovery: a collection run searches a chosen geography and selected business categories through the configured business-data provider.', bullet_id)
    add_bullet(doc, 'CSV import: a user supplies a prepared client list for fast seeding, testing, or sources not connected to live discovery.', bullet_id)
    add_bullet(doc, 'Existing workspace client: a client can be re-analyzed or enriched without being newly discovered.', bullet_id)
    add_heading(doc, 'Business information to collect when available', 2)
    add_para(doc, 'The client profile should make the source and freshness of each field understandable. Provider-supplied data can include:')
    for text in [
        'Business identity: name, category, address/location, website, source URL, phone, public email, rating, review count, hours, and service descriptions.',
        'Google/business profile signals: claimed status, review distribution and recency, photos, booking or estimate links, categories, and observed locations.',
        'Company metrics when a named source supplies them: employee count, estimated annual revenue, estimated annual profit, founding year, and branch count.',
        'Website and technology signals: CMS/framework clues, analytics or third-party scripts, booking/form systems, HTTPS, mobile setup, and other observable technical evidence.',
    ]:
        add_bullet(doc, text, bullet_id)
    add_note(doc, 'Unknowns are valid data', 'The product must not infer revenue, profit, employee count, branch count, decision-maker identity, hosting cost, or business age from weak signals. When no source supplies a field, show it as unknown and explain what source would be required.')
    add_heading(doc, 'Run identity and provenance', 2)
    add_para(doc, 'Each run is identifiable by an ordered run number, start time, source, location, categories, status, and progress. The run should expose the clients it touched and distinguish:')
    add_bullet(doc, 'New clients: first added to the workspace by that run.', bullet_id)
    add_bullet(doc, 'Matched clients: already present but returned again by that run.', bullet_id)
    add_bullet(doc, 'Audited clients: clients whose website research was processed by that run.', bullet_id)
    add_para(doc, 'Each client should retain a human-readable origin such as “Added in Run 1” and a history of later runs that matched or audited it. This avoids the ambiguity of a single undated “source” label.')

    add_heading(doc, '7. Website analysis requirements', 1)
    add_para(doc, 'Website analysis should be bounded, repeatable, and honest about partial results. It is intended to reveal actionable business and conversion opportunities, not to produce a generic technical audit report.')
    add_heading(doc, 'Initial observations', 2)
    for text in [
        'Reachability, redirect behavior, HTTPS, and basic response success.',
        'Page title, meta description, canonical, viewport, heading hierarchy, structured data, and content signals.',
        'Contact paths, phone/email links, forms, CTA visibility, quote/booking paths, financing language, and trust signals.',
        'Image alt coverage, mobile setup, technology clues, scripts, and other observable implementation signals.',
    ]:
        add_bullet(doc, text, bullet_id)
    add_heading(doc, 'Bounded crawl and link checks', 2)
    add_bullet(doc, 'Use the homepage plus a limited number of same-site pages, optionally informed by the sitemap.', bullet_id)
    add_bullet(doc, 'Respect the website’s robots rules where applicable and avoid action routes such as login, checkout, or account paths.', bullet_id)
    add_bullet(doc, 'Check a bounded set of internal links and save the response status or error for each checked destination.', bullet_id)
    add_bullet(doc, 'Record crawl errors and partial completion rather than hiding failed pages.', bullet_id)
    add_heading(doc, 'Lighthouse and measured performance', 2)
    add_bullet(doc, 'Where available, collect mobile performance, accessibility, SEO, best-practice, and timing metrics.', bullet_id)
    add_bullet(doc, 'If Lighthouse cannot launch or load the page, keep the client analyzed with a visible partial component state; never fabricate a performance score.', bullet_id)
    add_heading(doc, 'What the analysis should produce', 2)
    add_para(doc, 'The result should include the measured score and confidence, a prioritized set of problems, supporting evidence, crawl/page/link counts, technical metrics, and a clear indication of which parts are complete, partial, unavailable, or failed.')

    add_heading(doc, '8. Scoring and recommendation logic', 1)
    add_para(doc, 'The opportunity score is a prioritization aid for deciding which clients deserve attention. It is not a prediction of revenue, a ranking checkpoint, or an automatic qualification decision.')
    add_heading(doc, 'Score behavior', 2)
    add_bullet(doc, 'Use measured business fit, website/conversion gaps, technical pain, rebuild feasibility, and contactability signals where available.', bullet_id)
    add_bullet(doc, 'Keep confidence separate from opportunity. An exciting lead with thin evidence should be clearly marked as uncertain.', bullet_id)
    add_bullet(doc, 'Every scored issue should be explainable through an evidence point, metric, URL, or saved observation.', bullet_id)
    add_bullet(doc, 'Muse may prioritize or explain findings but must not silently change the deterministic score.', bullet_id)
    add_heading(doc, 'Recommendation semantics', 2)
    add_para(doc, '“Recommended” means “ready for a human decision based on the available evidence.” It is not the same as “qualified,” “contacted,” or “won.” The user must explicitly qualify a client before it enters the pipeline.')

    add_heading(doc, '9. Muse synthesis requirements', 1)
    add_para(doc, 'Muse is the optional reasoning layer that turns saved evidence into a compact decision brief. It should be treated as a collaborator that organizes evidence, not as the authority that invents facts or owns the score.')
    add_heading(doc, 'Muse inputs', 2)
    add_bullet(doc, 'The saved client profile and source-backed business signals.', bullet_id)
    add_bullet(doc, 'The deterministic opportunity score and confidence.', bullet_id)
    add_bullet(doc, 'The website crawl summary, page observations, broken-link evidence, and Lighthouse metrics or selected failures.', bullet_id)
    add_bullet(doc, 'The user’s configured outreach-length preference when drafting outreach.', bullet_id)
    add_heading(doc, 'Muse outputs', 2)
    for text in [
        'A concise summary of the business and the most important opportunity.',
        'Prioritized opportunities with severity, observed evidence, and a recommended improvement.',
        'Strengths worth preserving or mentioning.',
        'Unknowns and evidence gaps that should not be presented as facts.',
        'Qualification guidance such as pursue, review, or deprioritize, with rationale.',
        'Outreach angles that are grounded in the actual findings.',
    ]:
        add_bullet(doc, text, bullet_id)
    add_note(doc, 'No extra checkpoint', 'Muse does not create a separate “LLM insights” stage. The user should experience it as an optional synthesis inside the analyzed/recommendation context, while the underlying website analysis remains useful without it.')

    add_heading(doc, '10. Screen-by-screen functional requirements', 1)
    add_heading(doc, '10.1 Clients — the client database', 2)
    add_para(doc, 'Purpose: provide a reliable workspace-level view of every business and its current research and sales state.')
    add_bullet(doc, 'Show aggregate counts for total clients, qualified clients, open pipeline clients, and won clients.', bullet_id)
    add_bullet(doc, 'Support views for All, Researching, In pipeline, and Won.', bullet_id)
    add_bullet(doc, 'Allow client search by name and related profile text.', bullet_id)
    add_bullet(doc, 'Show each client’s research phase, opportunity score, sales stage, last update, and origin run.', bullet_id)
    add_bullet(doc, 'Open a client’s detail workspace without losing the current navigation context.', bullet_id)
    add_heading(doc, '10.2 Research — the evidence and operations workspace', 2)
    add_para(doc, 'Purpose: make research progress legible from raw collection to a human qualification decision.')
    add_bullet(doc, 'Show counts for Scraped, Analyzed, and Recommended clients.', bullet_id)
    add_bullet(doc, 'Show each client’s three research checkpoints, current phase, failure state, score, and next available action.', bullet_id)
    add_bullet(doc, 'Allow a user to analyze one client, retry a failed audit, or qualify a completed client.', bullet_id)
    add_bullet(doc, 'Show the currently active collection run with its true phase, status, progress, and controls.', bullet_id)
    add_bullet(doc, 'Show run history, including run number/date/source, new and matched client counts, audited count, and the clients linked to that run.', bullet_id)
    add_heading(doc, '10.3 Client detail workspace', 2)
    add_para(doc, 'Purpose: provide a complete, evidence-backed view of one business and all actions that can be taken on it.')
    add_bullet(doc, 'Business identity, category, address, website, Google Maps link, contact paths, and source context.', bullet_id)
    add_bullet(doc, 'Origin run and later run history, including whether the client was new, matched, or audited.', bullet_id)
    add_bullet(doc, 'Research progress across Scraped, Analyzed, and Recommended.', bullet_id)
    add_bullet(doc, 'Summary of the main opportunity, findings, business intelligence, score, confidence, crawl metrics, Lighthouse metrics, and unknowns.', bullet_id)
    add_bullet(doc, 'Manual notes and editable operational context.', bullet_id)
    add_bullet(doc, 'Manual actions: refresh business data, run website audit, run Muse analysis, generate outreach, or run the full workflow.', bullet_id)
    add_bullet(doc, 'Sales tab with relationship stage, outreach draft, and contract details after a win.', bullet_id)
    add_heading(doc, '10.4 Outreach — the review-before-send workspace', 2)
    add_para(doc, 'Purpose: turn qualified research into a short, relevant, human-reviewed first message.')
    add_bullet(doc, 'Show qualified clients with the strongest relevant findings alongside their draft.', bullet_id)
    add_bullet(doc, 'Support All, Ready, and Contacted views.', bullet_id)
    add_bullet(doc, 'Generate, edit, save, select, and copy one of three email versions.', bullet_id)
    add_bullet(doc, 'Show whether an email was generated from the configured model or a deterministic evidence fallback.', bullet_id)
    add_bullet(doc, 'Allow the user to mark a client Contacted only after reviewing the draft; sending itself remains outside the app.', bullet_id)
    add_heading(doc, '10.5 Pipeline — the relationship board', 2)
    add_para(doc, 'Purpose: manage the manual sales relationship after a client is qualified.')
    add_bullet(doc, 'Use the stages Qualified, Contacted, Replied, Meeting, Proposal, Won, and Lost.', bullet_id)
    add_bullet(doc, 'Allow moving a client by drag-and-drop or an explicit stage menu.', bullet_id)
    add_bullet(doc, 'Treat Qualified as the entry point to the pipeline, not as a synonym for research recommendation.', bullet_id)
    add_bullet(doc, 'Keep stage changes visible on the client record and in the relevant navigation counts.', bullet_id)
    add_heading(doc, '10.6 Won — the handoff and contract workspace', 2)
    add_para(doc, 'Purpose: record the agreement once a client says yes and preserve the information needed for delivery follow-up.')
    add_bullet(doc, 'Show won-client count, recorded contract value, and start-date coverage.', bullet_id)
    add_bullet(doc, 'Capture approved amount, currency, scope, payment terms, signed date, start date, conditions, and next step.', bullet_id)
    add_bullet(doc, 'Allow contract details to be added or edited manually.', bullet_id)
    add_heading(doc, '10.7 Settings and workspace state', 2)
    add_para(doc, 'Purpose: expose only the defaults and provider readiness needed for day-to-day operation.')
    add_bullet(doc, 'Default location, client limit, deep-audit limit, and discovery coordinates.', bullet_id)
    add_bullet(doc, 'Whether live business discovery is available.', bullet_id)
    add_bullet(doc, 'Which outreach model/provider is configured and whether a fallback is active.', bullet_id)
    add_bullet(doc, 'Local workspace and persistence status at a conceptual level.', bullet_id)

    add_heading(doc, '11. Outreach content requirements', 1)
    add_para(doc, 'Outreach should be concise, specific, and easy for a busy business owner to understand. The system should use the evidence to choose one or two high-value observations rather than dumping the entire audit into a message.')
    outreach_rows = [
        ('Concise', 'Around 50 words', 'One or two important findings, one clear improvement idea, and a low-friction invitation.'),
        ('Balanced', 'Around 110 words', 'A short context line, a small set of evidence-backed findings, why they matter, and a simple offer.'),
        ('Detailed', 'Materially longer and more complete', 'A fuller report of the main issues, supporting evidence, possible improvements, and a clear next step. It should still be readable, not a raw audit dump.'),
    ]
    add_table(doc, ['Version', 'Length target', 'Functional requirement'], outreach_rows, [1450, 1800, 6210])
    add_heading(doc, 'Generation rules', 2)
    add_bullet(doc, 'The selected length preference must be included in the model instruction when drafts are generated.', bullet_id)
    add_bullet(doc, 'The three versions must be meaningfully different in detail, not the same message with a few words removed.', bullet_id)
    add_bullet(doc, 'Claims must come from saved evidence. If a metric is unknown, the email must not invent it.', bullet_id)
    add_bullet(doc, 'The draft should offer a relevant next step, such as sharing a free homepage concept or discussing one improvement.', bullet_id)
    add_bullet(doc, 'The user can edit any subject or body before copying it.', bullet_id)

    add_heading(doc, '12. Sales and contract requirements', 1)
    stage_rows = [
        ('Qualified', 'The user has chosen to pursue the client; the client enters the pipeline.'),
        ('Contacted', 'The user has reviewed and sent or otherwise delivered the first outreach.'),
        ('Replied', 'The client has responded.'),
        ('Meeting', 'A conversation or meeting is scheduled or has taken place.'),
        ('Proposal', 'A proposal or commercial scope has been shared.'),
        ('Won', 'The client agreed to work together; contract details can be recorded.'),
        ('Lost', 'The opportunity is closed without a win and should remain historically visible.'),
    ]
    add_table(doc, ['Sales stage', 'Meaning'], stage_rows, [1700, 7660])
    add_heading(doc, 'Stage rules', 2)
    add_bullet(doc, 'A client can be Recommended without being Qualified.', bullet_id)
    add_bullet(doc, 'Qualifying a client sets the sales stage to Qualified and makes it available in Outreach and Pipeline.', bullet_id)
    add_bullet(doc, 'Moving to any later sales stage implies the client is qualified.', bullet_id)
    add_bullet(doc, 'Won is a relationship outcome, not just a label; it unlocks the contract record and delivery handoff context.', bullet_id)

    add_heading(doc, '13. Manual controls and individual-client workflows', 1)
    add_para(doc, 'The user should not need to rerun a full batch to fix or refresh one client. Every client can be operated on individually when the user has a specific reason.')
    manual_rows = [
        ('Refresh business data', 'Re-fetch the provider record for a client with a saved source URL.', 'Bright Data/provider data is refreshed; missing fields can be enriched.'),
        ('Run website audit', 'Repeat the bounded website analysis for one client.', 'Research evidence and score are refreshed.'),
        ('Run Muse analysis', 'Synthesize the saved audit evidence for one client.', 'Muse recommendation, opportunities, strengths, and unknowns are refreshed.'),
        ('Generate outreach', 'Create the three evidence-based drafts for one qualified or analyzed client.', 'Subject and drafts are saved for review.'),
        ('Run full workflow', 'Refresh business data, audit, synthesize, and generate outreach in sequence where prerequisites are available.', 'The client receives a visible multi-step manual run state.'),
    ]
    add_table(doc, ['Action', 'When to use it', 'Expected result'], manual_rows, [1900, 3500, 3960])
    add_para(doc, 'Each manual action needs visible Queued, Running, Waiting, Completed, and Failed states where relevant. A failed action should explain which phase failed, preserve the successful earlier work, and offer a retry instead of forcing the user to start over.')

    add_heading(doc, '14. Failure, partial-data, and recovery behavior', 1)
    add_para(doc, 'The product should make failure understandable without turning the workspace into an error console. Failures are operational facts about a run or action, not judgments about the prospect.')
    failure_rows = [
        ('Discovery/provider failure', 'The collection provider cannot start, returns an error, or never completes.', 'Mark the run as failed or waiting with the provider context; preserve the run record and allow retry.'),
        ('Website fetch failure', 'A site cannot be reached, resolves to a blocked/private target, times out, or returns an unusable response.', 'Mark the audit stage failed; keep the client record and source data; allow a retry or manual review.'),
        ('Partial crawl or Lighthouse failure', 'Some pages or a measurement component cannot be collected.', 'Keep the analyzed result with a visible partial component status and never substitute a fabricated metric.'),
        ('Muse failure', 'The model returns an error, invalid structure, or cannot be reached.', 'Keep the deterministic analysis and score; mark Muse failed and offer retry.'),
        ('Outreach generation failure', 'The selected provider or fallback cannot produce a draft.', 'Keep findings available; explain the failure; allow retry or manual editing.'),
        ('Manual run failure', 'A multi-step client workflow stops after one phase.', 'Preserve completed phases, identify the failed phase, and let the user rerun that action.'),
    ]
    add_table(doc, ['Failure type', 'Meaning', 'Required behavior'], failure_rows, [1850, 3150, 4360])

    add_heading(doc, '15. Data integrity and trust requirements', 1)
    for text in [
        'Preserve source and observed-at context for important business and website fields.',
        'Keep repeated run history instead of overwriting the fact that a client was seen or audited before.',
        'Deduplicate repeated provider results without losing the client’s original origin run or later matched runs.',
        'Keep deterministic score, confidence, Muse synthesis, findings, outreach drafts, sales stage, and contract record conceptually separate.',
        'Do not let a model invent metrics, customer value, decision-makers, outcomes, or contract terms.',
        'Do not send outreach automatically, and design for human review and future suppression/opt-out handling.',
        'Treat website data as potentially hostile or incomplete input; preserve safe, bounded, and partial results.',
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, '16. Current scope versus intentionally future work', 1)
    add_heading(doc, 'In the current product scope', 2)
    for text in [
        'Local-first client workspace with persistent clients, runs, jobs, settings, and event history.',
        'Live business discovery through the configured provider or CSV import.',
        'Source-backed profile enrichment and explicit unknown values.',
        'Website HTML inspection, bounded crawl, internal-link checks, Lighthouse measurements, and saved evidence.',
        'Deterministic opportunity score and confidence.',
        'Optional Muse synthesis grounded in saved evidence.',
        'Three outreach draft lengths with edit/copy/review workflow.',
        'Manual qualification, Kanban-style sales stages, and contract recording after a win.',
        'Per-client manual actions, retryable job states, failures, and run provenance.',
    ]:
        add_bullet(doc, text, bullet_id)
    add_heading(doc, 'Not part of the current product behavior', 2)
    for text in [
        'Automatic cold-email sending, automatic follow-up sequences, or automatic qualification.',
        'A separate LLM-insights checkpoint or a separate ranked checkpoint.',
        'A complete competitor benchmarking system or longitudinal “why now?” monitoring.',
        'Automatic public redesign publishing or a public speculative website for every lead.',
        'Guaranteed employee, revenue, profit, branch, decision-maker, hosting-cost, or website-age data when no source provides it.',
        'A full CRM, billing system, team permissions, or multi-workspace collaboration model.',
        'A replacement for human sales judgment, compliance review, or contract/legal review.',
    ]:
        add_bullet(doc, text, bullet_id)
    add_heading(doc, 'Potential next-stage capabilities', 2)
    for text in [
        'Competitor comparisons that show relative performance, conversion friction, and service-page coverage.',
        'Revisit schedules and “why now?” signals such as new locations, new services, outages, or website changes.',
        'Small proof artifacts such as an annotated screenshot, homepage concept, or short opportunity report for manually approved leads.',
        'Decision-maker and contact enrichment from lawful, user-approved sources.',
        'Suppression lists, outreach limits, audit history comparison, and a more complete revenue learning loop.',
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, '17. Design questions for the next exploration', 1)
    questions = [
        'How should a user compare two runs without losing the simple current workflow?',
        'Should the default Research view show the latest run only, all active work, or a user-selected run?',
        'How much provenance belongs on a client list row versus the client detail workspace?',
        'What is the clearest distinction between “Recommended,” “Qualified,” and “In pipeline” for a first-time user?',
        'Should a client’s run history show only discovery events, or every audit and enrichment event?',
        'How should the product communicate partial evidence without making a client feel “broken”?',
        'What information is necessary before a user feels comfortable copying an outreach draft?',
        'Which proof artifact is most valuable for a qualified lead: a concise report, annotated screenshot, or private concept?',
        'What should the user see immediately after a run completes: new clients, highest opportunities, failures, or changes since the previous run?',
        'Which signals should be learned from real outcomes such as reply, meeting, proposal, and win?',
    ]
    for question in questions:
        add_numbered(doc, question, decimal_id)

    add_heading(doc, 'Appendix A — canonical vocabulary', 1)
    vocab_rows = [
        ('Research phases', 'Scraped · Analyzed · Recommended'),
        ('Operational statuses', 'Queued · Running · Waiting · Completed · Failed · Paused · Cancelled'),
        ('Sales stages', 'Qualified · Contacted · Replied · Meeting · Proposal · Won · Lost'),
        ('Run relationships', 'New · Matched · Audited'),
        ('Score concepts', 'Opportunity score · Confidence · Evidence · Unknowns'),
        ('AI concept', 'Muse synthesis — advisory, evidence-grounded, and separate from the deterministic score'),
    ]
    add_table(doc, ['Vocabulary group', 'Canonical terms'], vocab_rows, [2200, 7160])
    add_note(doc, 'Final product definition', 'Scoutline helps a human researcher discover promising local businesses, understand their website and business opportunities with evidence, choose who is worth pursuing, prepare a relevant outreach conversation, and manage the relationship through a manual sales pipeline and contract handoff.')

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build_document()
